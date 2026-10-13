"""Excel / CSV export: the measurements in one sheet (file), the average groups in another."""

import csv
import os
from datetime import datetime

from ...i18n import t

MEASURE_COLUMNS = ("n", "timestamp", "elapsed", "x", "y", "z", "position", "kind", "raw", "value", "unit", "group")
GROUP_COLUMNS = ("group", "first", "duration", "count", "x", "y", "z", "mean", "std", "min", "max", "unit")


def _stamp(seconds):
    return datetime.fromtimestamp(seconds).strftime("%Y-%m-%d %H:%M:%S.%f")[:-3]


def measurement_rows(measurements):
    """Rows of the first table, in the order of MEASURE_COLUMNS."""
    return [[m.n, _stamp(m.timestamp), round(m.elapsed_s, 6), m.x, m.y, m.z, t(f"results.position.{m.position_kind}"),
             t(f"results.kind.{m.kind}"), m.raw, m.value, m.unit, m.group_id] for m in measurements]


def group_rows(groups):
    """Rows of the second table, in the order of GROUP_COLUMNS."""
    return [[g.id, _stamp(g.first_timestamp) if g.count else None, round(g.duration_s, 6), g.count, g.x, g.y, g.z,
             g.mean, g.std, g.min, g.max, g.unit] for g in groups]


def _headers(columns):
    return [t(f"results.col.{c}") for c in columns]


def export_xlsx(store, path):
    """Write the workbook; returns the number of measurements. Needs openpyxl."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    measurements, groups = store.snapshot()
    book = Workbook()
    tables = ((t("results.sheet.measurements"), MEASURE_COLUMNS, measurement_rows(measurements)),
              (t("results.sheet.groups"), GROUP_COLUMNS, group_rows(groups)))
    for index, (title, columns, rows) in enumerate(tables):
        sheet = book.active if index == 0 else book.create_sheet()
        sheet.title = title
        sheet.append(_headers(columns))
        for row in rows:
            sheet.append(row)
        for cell in sheet[1]:
            cell.font, cell.alignment = Font(bold=True), Alignment(horizontal="center", wrap_text=True)
        sheet.freeze_panes = "A2"
        for number, header in enumerate(_headers(columns), start=1):
            width = max([len(header)] + [len(str(r[number - 1])) for r in rows[:200] if r[number - 1] is not None])
            sheet.column_dimensions[get_column_letter(number)].width = min(28, width + 3)
    book.save(path)
    return len(measurements)


def export_csv(store, path, delimiter=";", decimal=","):
    """Measurements in `path`, groups in `<name>_groups.csv`."""
    measurements, groups = store.snapshot()

    def cell(value):
        if value is None:
            return ""
        if isinstance(value, float):
            text = f"{value:.9g}"
            return text.replace(".", decimal) if decimal != "." else text
        return value

    def write(target, columns, rows):
        with open(target, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f, delimiter=delimiter)
            writer.writerow(_headers(columns))
            writer.writerows([[cell(v) for v in row] for row in rows])

    write(path, MEASURE_COLUMNS, measurement_rows(measurements))
    if groups:
        base, ext = os.path.splitext(path)
        write(f"{base}_groups{ext}", GROUP_COLUMNS, group_rows(groups))
    return len(measurements)
