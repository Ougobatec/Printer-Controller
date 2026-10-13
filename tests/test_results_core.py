import pytest

from printer_controller import i18n
from printer_controller.core.results import export_images
from printer_controller.core.results.export_xlsx import export_csv, export_xlsx
from printer_controller.core.results.store import ResultStore


def filled_store():
    store = ResultStore()
    store.start_run()
    group = store.new_group()
    for i in range(4):
        store.add(1000.0 + i * 0.1, {"X": 10, "Y": 20, "Z": 5}, "defined", "average", 0.001 * i, float(i), "mV", group)
    store.close_group(group)
    store.add(1002.0, {"X": 10, "Y": 20, "Z": 5}, "defined", "simple", 0.005, 5.0, "mV")
    for i in range(6):
        store.add(1003.0 + i * 0.1, {"X": 10 + i * 5, "Y": 20, "Z": 5}, "computed", "path", 0.01 * i, 10.0 + i, "mV")
    return store


def test_excel_has_two_sheets_with_all_values(tmp_path):
    from openpyxl import load_workbook
    i18n.set_language("en")
    path = tmp_path / "r.xlsx"
    assert export_xlsx(filled_store(), path) == 11
    book = load_workbook(path)
    assert book.sheetnames == ["Measurements", "Groups"]
    sheet = book["Measurements"]
    header = [c.value for c in sheet[1]]
    assert header[:3] == ["No.", "Timestamp", "Time since start (s)"] and header[-1] == "Group"
    assert sheet.max_row == 12
    first, last = [c.value for c in sheet[2]], [c.value for c in sheet[12]]
    assert first[6:8] == ["defined", "average"] and first[11] == 1
    assert last[6:8] == ["computed", "path"] and last[11] is None
    groups = book["Groups"]
    assert [c.value for c in groups[1]][7:11] == ["Mean", "Standard deviation", "Min", "Max"]
    row = [c.value for c in groups[2]]
    assert row[0] == 1 and row[3] == 4 and row[7] == pytest.approx(1.5) and row[10] == 3.0 and row[11] == "mV"


def test_csv_uses_the_chosen_separators(tmp_path):
    i18n.set_language("en")
    path = tmp_path / "r.csv"
    export_csv(filled_store(), path, delimiter=";", decimal=",")
    lines = path.read_text(encoding="utf-8-sig").splitlines()
    assert len(lines) == 12 and lines[0].startswith("No.;Timestamp")
    assert ";0,001;" in lines[2]
    assert (tmp_path / "r_groups.csv").exists()


def test_png_export(tmp_path):
    i18n.set_language("en")
    store = filled_store()
    measurements, _ = store.snapshot()
    for view in ("time", "x", "z"):
        path = tmp_path / f"{view}.png"
        shown, total = export_images.render_2d(path, measurements, view, "mV")
        assert path.stat().st_size > 5000 and total >= 1
    path = tmp_path / "3d.png"
    shown, total = export_images.render_3d(path, measurements, "mV", 45, 30)
    assert path.stat().st_size > 5000 and shown >= 1 and total == 11
    with pytest.raises(ValueError):
        export_images.render_3d(tmp_path / "x.png", [], "mV", 0, 0)


def test_view_angles_match_the_screen_presets():
    import math
    azim, elev = export_images.view_to_angles(0.68 + math.pi, 0.52)
    assert azim == pytest.approx(51, abs=1) and elev == pytest.approx(29.8, abs=0.5)
